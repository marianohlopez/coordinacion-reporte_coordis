from openpyxl import Workbook
from openpyxl.styles import Font
from dotenv import load_dotenv
import os
import yagmail
from extract import (extract_coordis, extract_pas, extract_informes, extract_seguim,
                     extract_seguim_mes)

load_dotenv()

MAIL_AUTOR = os.getenv("MAIL_AUTOR")
APP_GMAIL_PASS = os.getenv("APP_GMAIL_PASS")

def enviar_mail(destinatario, subject, body, archivo_adjunto):
  try:
    yag = yagmail.SMTP(MAIL_AUTOR, APP_GMAIL_PASS)
    yag.send(
        to=destinatario,
        subject=subject,
        contents=body,
        attachments=archivo_adjunto
    )
    print(f"📧 Mail enviado a {destinatario}")
  except Exception as e:
    print(f"❌ Error enviando mail a {destinatario}: {e}")


def generar_excel_por_coordinadora(coord_nombre, pas, informes, seguimientos, seguimientos_mes,
                                   coord_mail):
  wb = Workbook()
  ws = wb.active
  ws.title = "PAs disponibles"

  headers = ["PA ID", "NOMBRE", "ETIQUETA", "ESTADO", "OBSERVACIONES", "DISPONIBILIDAD",
              "REF. DE BUSQUEDA", "LOCALIDAD", "TELEFONO", "TELEFONO 2","EMAIL"]
  ws.append(headers)
  for cell in ws[1]:
      cell.font = Font(bold=True)

  for pa in pas:
      ws.append(pa)

# Segunda hoja (informes)
  ws2 = wb.create_sheet(title="Detalle de informes")

  headers_informes = ["COORDINADORA", "ALUMNO", "DNI ALUMNO", "INF. ADMISIÓN", "CONF. PA", "INF_SOCIAL",
                      "INF. MENSUAL", "INF. DIAGNÓSTICO", "INF. MEDIO", "PAC", "OTRO", "AA", "PPI", 
                      "INF. FINAL", "CONF. FLIA.", "INF. ESCOLAR", "INF. TER. EXT.", 
                      "PLAN TRAB. COORD."]
  
  ws2.append(headers_informes)

  for cell in ws2[1]:
    cell.font = Font(bold=True)

  for row in informes:
      ws2.append(row)

  # Tercer hoja (seguimientos_detalle)
  ws3 = wb.create_sheet(title="Detalle de seguimientos")

  headers_seguimientos = ["USUARIO DE CARGA", "ROL", "PRESTACION ID", "TIPO", "ALUMNO", 
                          "MES", "FECHA DE CARGA", "CATEG. SEGUIMIENTO"]
  
  ws3.append(headers_seguimientos)

  for cell in ws3[1]:
    cell.font = Font(bold=True)

  for row in seguimientos:
    ws3.append(row)

  # Cuarta hoja (seguimientos_mes)
  ws4 = wb.create_sheet(title="Seguimientos por mes")

  headers_seguimientos_mes = ["ALUMNO", "PRESTACION ID", "TIPO", "USUARIO DE CARGA", "ROL", "ENE", 
                              "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", 
                              "NOV", "DIC", "TOTAL ANUAL"]
  
  ws4.append(headers_seguimientos_mes)

  for cell in ws4[1]:
    cell.font = Font(bold=True)

  for row in seguimientos_mes:
    ws4.append(row)

  filename = f"reporte_{coord_nombre.replace(', ', '_')}.xlsx"
  wb.save(filename)
  print(f"✅ Excel generado: {filename}")

  enviar_mail(
    destinatario=coord_mail,
    subject=f"Reporte de PAs disponibles - {coord_nombre} (NO CONTESTAR)",
    body=f"""Hola {coord_nombre},\n\nSe adjunta el listado actualizado de PAs disponibles 
    en sus localidades junto con el detalle de los informes y seguimientos
    cargados en Indyco.\n\nSaludos,\nMariano López - Ailes Inclusión.""",
    archivo_adjunto=filename
  )

def generar_reportes_por_coordinadora(conn):
  cursor = conn.cursor()

  cant = 0
  registros = 0

  coordinadoras = extract_coordis(cursor)

  for coord in coordinadoras:
    coord_id = coord[0]
    coord_nombre = coord[1]
    coord_mail = coord[2]
    coord_user_id = coord[3]

    pas = extract_pas(cursor, coord_id)
    informes = extract_informes(cursor, coord_id)
    seguimientos = extract_seguim(cursor, coord_id)
    seguimientos_mes = extract_seguim_mes(cursor, coord_user_id)
    generar_excel_por_coordinadora(coord_nombre, pas, informes, seguimientos, seguimientos_mes,
                                   coord_mail)
    cant += 1
    registros += len(pas)

  return cant, registros